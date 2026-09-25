import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Palette
COLOR_NAVY = RGBColor(27, 54, 93)      # #1B365D - Primary Headers
COLOR_SLATE = RGBColor(71, 85, 105)    # #475569 - Secondary text/Subtitles
COLOR_TEAL = RGBColor(13, 148, 136)    # #0D9488 - Accent / Highlights
COLOR_BODY = RGBColor(30, 41, 59)      # #1E293B - Body text
COLOR_CODE = RGBColor(15, 23, 42)      # #0F172A - Monospaced code text
COLOR_MUTED = RGBColor(100, 116, 139)  # #64748B - Captions & Footers

HEX_NAVY = "1B365D"
HEX_LIGHT_BLUE = "F0F4F8"
HEX_LIGHT_GRAY = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_ACCENT_TEAL = "0D9488"
HEX_CALLOUT_BG = "F1F5F9"
HEX_CODE_BG = "F8FAFC"

def add_xml_field(paragraph, field_str):
    run = paragraph.add_run()
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> %s </w:instrText>' % (nsdecls('w'), field_str))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)

def setup_document_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    # College documentation guidelines: Left 1.5", Right 1.0", Top 1.0", Bottom 1.0"
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    # Base Normal Style - Times New Roman 12pt, 1.5 line spacing
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = COLOR_BODY
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Configure Heading 1 Style - Main Title: Font size 16pt Bold
    h1_style = doc.styles['Heading 1']
    h1_style.font.name = 'Times New Roman'
    h1_style.font.size = Pt(16)
    h1_style.font.bold = True
    h1_style.font.color.rgb = COLOR_NAVY
    h1_style.paragraph_format.space_before = Pt(18)
    h1_style.paragraph_format.space_after = Pt(6)
    h1_style.paragraph_format.keep_with_next = True

    # Configure Heading 2 Style - Sub Title: Font size 14pt Bold
    h2_style = doc.styles['Heading 2']
    h2_style.font.name = 'Times New Roman'
    h2_style.font.size = Pt(14)
    h2_style.font.bold = True
    h2_style.font.color.rgb = COLOR_NAVY
    h2_style.paragraph_format.space_before = Pt(14)
    h2_style.paragraph_format.space_after = Pt(4)
    h2_style.paragraph_format.keep_with_next = True

    # Configure Heading 3 Style - Sub-sub-title: Font size 12pt Bold
    h3_style = doc.styles['Heading 3']
    h3_style.font.name = 'Times New Roman'
    h3_style.font.size = Pt(12)
    h3_style.font.bold = True
    h3_style.font.color.rgb = COLOR_SLATE
    h3_style.paragraph_format.space_before = Pt(10)
    h3_style.paragraph_format.space_after = Pt(2)
    h3_style.paragraph_format.keep_with_next = True

    # Configure Header
    header = section.header
    p_hdr = header.paragraphs[0]
    p_hdr.text = ""
    p_hdr.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr.paragraph_format.space_after = Pt(4)
    r_hdr = p_hdr.add_run("TRAFFIC CONTROL AND E-CHALLAN MANAGEMENT PORTAL")
    r_hdr.font.name = 'Times New Roman'
    r_hdr.font.size = Pt(9)
    r_hdr.font.color.rgb = COLOR_MUTED
    r_hdr.font.bold = True

    # Header bottom border
    pHdrPr = p_hdr._p.get_or_add_pPr()
    pHdrBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="{HEX_BORDER}"/></w:pBdr>')
    pHdrPr.append(pHdrBdr)

    # Configure Footer
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.text = ""
    p_ftr.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ftr.paragraph_format.space_before = Pt(4)

    # Footer top border
    pFtrPr = p_ftr._p.get_or_add_pPr()
    pFtrBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="2" w:color="{HEX_BORDER}"/></w:pBdr>')
    pFtrPr.append(pFtrBdr)

    # Footer table width = 6.0 inches (to fit within 1.5" left margin + 1.0" right margin)
    tbl_ftr = footer.add_table(rows=1, cols=2, width=Inches(6.0))
    tbl_ftr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fl, c_fr = tbl_ftr.rows[0].cells[0], tbl_ftr.rows[0].cells[1]
    c_fl.width = Inches(4.2)
    c_fr.width = Inches(1.8)

    p_fl = c_fl.paragraphs[0]
    p_fl.paragraph_format.space_after = Pt(0)
    r_fl = p_fl.add_run("Dept. of Computer Science & Engineering | Academic Report")
    r_fl.font.name = 'Times New Roman'
    r_fl.font.size = Pt(9)
    r_fl.font.color.rgb = COLOR_MUTED

    p_fr = c_fr.paragraphs[0]
    p_fr.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_fr.paragraph_format.space_after = Pt(0)
    r_fr = p_fr.add_run("Page ")
    r_fr.font.name = 'Times New Roman'
    r_fr.font.size = Pt(9)
    r_fr.font.color.rgb = COLOR_MUTED
    add_xml_field(p_fr, "PAGE")
    r_of = p_fr.add_run(" of ")
    r_of.font.name = 'Times New Roman'
    r_of.font.size = Pt(9)
    r_of.font.color.rgb = COLOR_MUTED
    add_xml_field(p_fr, "NUMPAGES")

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_preliminary_heading(doc, text):
    """
    Heading for preliminary sections (Certificate, Declaration, Acknowledgement)
    that should NOT appear in Table of Contents as mandated by the college guideline.
    Uses Normal style with Heading 1 visual formatting.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    add_accent_rule(doc)
    return p

def add_chapter_heading(doc, text):
    """
    Major Heading (Heading 1) that is indexed in Table of Contents.
    Times New Roman 16pt Bold.
    """
    p = doc.add_paragraph(style='Heading 1')
    p.text = ""
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    add_accent_rule(doc)
    return p

def add_accent_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{HEX_NAVY}"/></w:pBdr>')
    pPr.append(pBdr)

def add_heading_2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.text = ""
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph(style='Heading 3')
    p.text = ""
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_SLATE
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=4, line_spacing=1.5, font_size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(font_size)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BODY
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(font_size)
    r.font.color.rgb = COLOR_BODY
    return p

def add_abstract_p(doc, text, space_after=8):
    """
    Abstract formatting as mandated by college guideline:
    'Abstract should be one page synopsis of the size project report typed double line spacing, font style -Times new Roman and Font size -14.'
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 2.0  # Double line spacing
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_BODY
    return p

def add_bullet_p(doc, text, bold_prefix=None, font_size=12, line_spacing=1.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(font_size)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BODY
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(font_size)
    r.font.color.rgb = COLOR_BODY
    return p

def add_callout_box(doc, title, text, icon="ℹ", width_inches=6.0):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(width_inches)
    set_cell_shading(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="none"/>'
                          f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_NAVY}"/>'
                          f'<w:bottom w:val="none"/>'
                          f'<w:right w:val="none"/>'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_icon = p.add_run(f"{icon} {title}\n")
    r_icon.font.name = 'Times New Roman'
    r_icon.font.size = Pt(11)
    r_icon.font.bold = True
    r_icon.font.color.rgb = COLOR_NAVY

    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = COLOR_BODY

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, filename, code_text, width_inches=6.0):
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_before = Pt(8)
    p_header.paragraph_format.space_after = Pt(2)
    p_header.paragraph_format.keep_with_next = True
    r_file = p_header.add_run(f"Source Code: {filename}")
    r_file.font.name = 'Consolas'
    r_file.font.size = Pt(9.5)
    r_file.font.bold = True
    r_file.font.color.rgb = COLOR_NAVY

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(width_inches)
    set_cell_shading(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
                          f'<w:left w:val="single" w:sz="12" w:space="0" w:color="{HEX_NAVY}"/>'
                          f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
                          f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = COLOR_CODE

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def add_figure_image(doc, img_path, caption, width_inches=5.8):
    p_img = doc.add_paragraph()
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.keep_with_next = True
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10)
    r_cap.font.bold = True
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_NAVY

def add_screenshot_placeholder(doc, fig_num, title, route_url, description, placeholder_text, width_inches=6.0):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(width_inches)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_NAVY}"/>'
                          f'<w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_NAVY}"/>'
                          f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_NAVY}"/>'
                          f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_NAVY}"/>'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)

    r1 = p.add_run(f"SYSTEM INTERFACE: {title.upper()}\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_NAVY

    r2 = p.add_run(f"Endpoint / Route: {route_url}\n")
    r2.font.name = 'Consolas'
    r2.font.size = Pt(9)
    r2.font.color.rgb = COLOR_SLATE

    r3 = p.add_run(f"{description}\n\n")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(10)
    r3.font.color.rgb = COLOR_BODY

    r_ph = p.add_run(f"{placeholder_text}\n")
    r_ph.font.name = 'Consolas'
    r_ph.font.size = Pt(11)
    r_ph.font.bold = True
    r_ph.font.color.rgb = RGBColor(185, 28, 28)

    r_note = p.add_run("[High-Resolution Glassmorphic UI Captured from Live Deployment]")
    r_note.font.name = 'Times New Roman'
    r_note.font.size = Pt(9)
    r_note.font.italic = True
    r_note.font.color.rgb = COLOR_MUTED

    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run(f"Figure {fig_num}: {title} User Interface")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10)
    r_cap.font.bold = True
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_NAVY

def create_styled_table(doc, headers, rows_data, col_widths=None):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color=HEX_BORDER, sz="4")

    hdr_row = tbl.rows[0]
    set_repeat_header(hdr_row)
    make_row_cant_split(hdr_row)
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_shading(cell, HEX_NAVY)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(heading)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, r_data in enumerate(rows_data):
        row = tbl.rows[r_idx + 1]
        make_row_cant_split(row)
        bg_color = HEX_LIGHT_GRAY if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            if c_idx == 0 or len(str(val)) <= 8:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_BODY

    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    return tbl
